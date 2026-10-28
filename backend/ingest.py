"""Turn an uploaded file into searchable chunks: [{page, chunk_index, text}].

Text formats are parsed locally. Scans, photos and recordings are read by the
generator's `transcribe(bytes, mime)` (Gemini in production), so a scanned PDF, a
phone photo of a notice or a voice memo all become searchable text.

Nothing readable should be lost to how a file was made: every PDF page is checked
for images (scans, handwriting, photos, diagrams), math or symbol fonts whose text
extracts badly, garbled text and near-empty text, and each such page is read by
the transcriber on its own. Form fields and comments are kept; images inside Word
and PowerPoint files are read too.
"""
from __future__ import annotations

import csv
import io
import ipaddress
import logging
import re
import socket
import urllib.parse
from concurrent.futures import ThreadPoolExecutor
from html.parser import HTMLParser
from typing import Callable

from fastapi import HTTPException

log = logging.getLogger("ragaas")

PDF_MAGIC = b"%PDF"
ZIP_MAGIC = b"PK\x03\x04"
# Gemini reads inline files up to ~20 MB per request.
MAX_TRANSCRIBE_BYTES = 19 * 1024 * 1024
ROWS_PER_CHUNK = 25

# Page-level OCR for PDFs (see _pdf_page_reasons).
MIN_PAGE_CHARS = 40          # less selectable text than this: likely a scan or an image page
MIN_IMAGE_SIDE = 64          # images smaller than this (px) are bullets and rules, not content
MAX_OCR_PAGES = 300          # per document; beyond this, local text only (logged)
OCR_WORKERS = 8
# Fonts whose glyphs extract out of order or without superscripts (10^9 → "10 … 9").
MATH_FONT = re.compile(r"KaTeX|MathJax|CMMI|CMSY|CMEX|MSBM|MSAM|CambriaMath|Cambria-Math|STIX|XITS|"
                       r"LatinModernMath|lmmath|Asana|TeXGyre\w*Math|Euclid|MT-?Extra", re.I)
# Embedded images in Word / PowerPoint worth reading (Gemini's supported image types).
EMBEDDED_IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/heic", "image/heif"}
MIN_EMBEDDED_IMAGE_BYTES = 4 * 1024
MAX_EMBEDDED_IMAGES = 40

IMAGE_TYPES = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "webp": "image/webp",
               "heic": "image/heic", "heif": "image/heif"}
AUDIO_TYPES = {"mp3": "audio/mpeg", "wav": "audio/wav", "m4a": "audio/mp4", "aac": "audio/aac",
               "ogg": "audio/ogg", "opus": "audio/ogg", "flac": "audio/flac", "webm": "audio/webm"}
TEXT_TYPES = {"txt", "md", "markdown"}
HTML_TYPES = {"html", "htm"}
SUPPORTED = (["pdf", "docx", "xlsx", "csv", "pptx"] + sorted(TEXT_TYPES | HTML_TYPES)
             + sorted(IMAGE_TYPES) + sorted(AUDIO_TYPES))
SUPPORTED_LABEL = "PDF, Word, Excel, CSV, PowerPoint, text, Markdown, HTML, photos (JPG/PNG) or audio (MP3/M4A/WAV)"

Transcriber = Callable[[bytes, str], str] | None


def chunk_text(text: str, chunk_size: int = 300, overlap: int = 50) -> list[str]:
    words = text.split()
    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = min(start + chunk_size, len(words))
        chunks.append(" ".join(words[start:end]))
        if end == len(words):
            break
        start += chunk_size - overlap
    return chunks or [""]


def _pages_to_chunks(pages: list[tuple[int, str]]) -> list[dict]:
    out = []
    for page, text in pages:
        for idx, chunk in enumerate(chunk_text(text)):
            if chunk.strip():
                out.append({"page": page, "chunk_index": idx, "text": chunk})
    return out


def _need(result: list[dict], what: str) -> list[dict]:
    if not result:
        raise HTTPException(status_code=422, detail=f"No readable text found in this {what}.")
    return result


def _ext(file_name: str) -> str:
    return file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""


# ── Transcription (scans, photos, audio) ─────────────────────────────────────
PAGE_MARK = re.compile(r"^\s*=+\s*page\s+(\d+)\s*=+\s*$", re.I | re.M)


def _transcribe(content: bytes, mime: str, transcribe: Transcriber, what: str) -> str:
    if transcribe is None:
        raise HTTPException(status_code=422, detail=f"Reading {what} needs the AI reader, which isn't available here.")
    if len(content) > MAX_TRANSCRIBE_BYTES:
        raise HTTPException(status_code=413, detail=f"{what.capitalize()} files can be up to 19 MB.")
    try:
        return transcribe(content, mime)
    except HTTPException:
        raise
    except Exception as exc:
        log.warning("transcribe failed mime=%s: %s", mime, exc)
        raise HTTPException(status_code=502, detail=f"Couldn't read this {what}. Try again, or upload a clearer file.")


def _clean_transcript(text: str) -> str:
    return PAGE_MARK.sub("", text or "").strip()


def _read_images(images: list[tuple[bytes, str]], transcribe: Transcriber) -> list[str]:
    """Text read from embedded images (Word/PowerPoint); unreadable ones are skipped."""
    if transcribe is None:
        return []
    picked = [(b, m) for b, m in images
              if m in EMBEDDED_IMAGE_TYPES and MIN_EMBEDDED_IMAGE_BYTES <= len(b) <= MAX_TRANSCRIBE_BYTES]
    if len(picked) > MAX_EMBEDDED_IMAGES:
        log.warning("reading %d of %d embedded images", MAX_EMBEDDED_IMAGES, len(picked))
        picked = picked[:MAX_EMBEDDED_IMAGES]

    def one(item: tuple[bytes, str]) -> str:
        try:
            return _clean_transcript(transcribe(item[0], item[1]))
        except Exception as exc:
            log.warning("embedded image unreadable (%s): %s", item[1], exc)
            return ""

    with ThreadPoolExecutor(max_workers=OCR_WORKERS) as pool:
        return [t for t in pool.map(one, picked) if t]


def _scanned_pdf(content: bytes, transcribe: Transcriber) -> list[dict]:
    text = _transcribe(content, "application/pdf", transcribe, "scanned PDF")
    # The reader is asked to mark pages as "=== Page N ===".
    parts = PAGE_MARK.split(text)
    if len(parts) > 1:
        pages = [(int(parts[i]), parts[i + 1]) for i in range(1, len(parts) - 1, 2)]
    else:
        pages = [(1, text)]
    return _need(_pages_to_chunks(pages), "scanned PDF")


# ── Local parsers ─────────────────────────────────────────────────────────────
def _images_on(resources, depth: int = 0) -> int:
    """Count content-sized images on a page, including inside form XObjects."""
    try:
        xobjects = (resources or {}).get("/XObject")
        xobjects = xobjects.get_object() if xobjects is not None else {}
    except Exception:
        return 0
    n = 0
    for ref in xobjects.values():
        try:
            x = ref.get_object()
            sub = x.get("/Subtype")
            if sub == "/Image":
                if int(x.get("/Width", 0)) >= MIN_IMAGE_SIDE and int(x.get("/Height", 0)) >= MIN_IMAGE_SIDE:
                    n += 1
            elif sub == "/Form" and depth < 3:
                n += _images_on(x.get("/Resources"), depth + 1)
        except Exception:
            continue
    return n


def _fonts_on(resources) -> list[str]:
    try:
        fonts = (resources or {}).get("/Font")
        fonts = fonts.get_object() if fonts is not None else {}
        return [str(f.get_object().get("/BaseFont", "")) for f in fonts.values()]
    except Exception:
        return []


def _pdf_page_reasons(page, text: str) -> list[str]:
    """Why a page's selectable text may not be everything on it (empty = it is)."""
    reasons = []
    try:
        resources = page.get("/Resources")
        resources = resources.get_object() if resources is not None else {}
    except Exception:
        resources = {}
    if len(text) < MIN_PAGE_CHARS:
        reasons.append("little or no selectable text")
    if _images_on(resources):
        reasons.append("images")
    if any(MATH_FONT.search(f) for f in _fonts_on(resources)):
        reasons.append("math or symbol fonts")
    if text and (text.count("\ufffd") + len(re.findall(r"\(cid:\d+\)", text))) > max(3, len(text) // 50):
        reasons.append("garbled text")
    return reasons


def _pdf_extras(page) -> str:
    """Filled-in form fields and comments: pypdf's text extraction skips both."""
    out = []
    try:
        annots = page.get("/Annots") or []
        for ref in annots.get_object() if hasattr(annots, "get_object") else annots:
            a = ref.get_object()
            value, name = a.get("/V"), a.get("/T")
            if value not in (None, "") and not hasattr(value, "get_object"):
                value = str(value).lstrip("/")
                if value and value not in ("Off",):
                    out.append(f"{name}: {value}" if name else value)
            contents = a.get("/Contents")
            if contents and a.get("/Subtype") in ("/Text", "/FreeText", "/Highlight", "/Underline", "/StrikeOut", "/Square", "/Circle"):
                out.append(f"Comment: {contents}")
    except Exception as exc:
        log.info("could not read annotations: %s", exc)
    return "\n".join(out)


def _one_page_pdf(reader, index: int) -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_page(reader.pages[index])
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def _pdf(content: bytes, transcribe: Transcriber) -> list[dict]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(content))
        texts = [(p.extract_text() or "").strip() for p in reader.pages]
        extras = [_pdf_extras(p) for p in reader.pages]
        reasons = [_pdf_page_reasons(p, t) for p, t in zip(reader.pages, texts)]
    except Exception as exc:
        log.warning("PDF parse failed: %s", exc)
        raise HTTPException(status_code=400, detail="Could not parse PDF: the file may be corrupted or encrypted.")

    flagged = [i for i, r in enumerate(reasons) if r]
    if flagged and transcribe is None and not any(texts):
        _transcribe(content, "application/pdf", None, "scanned PDF")      # raises the "needs the AI reader" error
    read: dict[int, str] = {}
    if flagged and transcribe is not None:
        if len(flagged) > MAX_OCR_PAGES:
            log.warning("PDF has %d pages to read; reading the first %d", len(flagged), MAX_OCR_PAGES)
            flagged = flagged[:MAX_OCR_PAGES]

        def one(i: int) -> tuple[int, str]:
            try:
                page_pdf = _one_page_pdf(reader, i)
                if len(page_pdf) > MAX_TRANSCRIBE_BYTES:
                    log.warning("PDF page %d is too large to read (%d bytes)", i + 1, len(page_pdf))
                    return i, ""
                return i, _clean_transcript(transcribe(page_pdf, "application/pdf"))
            except Exception as exc:
                log.warning("PDF page %d unreadable (%s): %s", i + 1, ", ".join(reasons[i]), exc)
                return i, ""

        with ThreadPoolExecutor(max_workers=OCR_WORKERS) as pool:
            read = {i: t for i, t in pool.map(one, flagged) if t}
        log.info("PDF pages read by the AI reader: %d of %d (%s)", len(read), len(reader.pages),
                 "; ".join(f"p{i + 1}: {', '.join(reasons[i])}" for i in flagged[:12]))

    pages = []
    for i, local in enumerate(texts):
        ocr = read.get(i, "")
        if ocr and len(ocr) >= 0.6 * len(local):
            text = ocr                       # the reader saw everything the text layer has, and more
        elif ocr:
            text = f"{local}\n\n[Also on this page]\n{ocr}"
        else:
            text = local
        if extras[i]:
            text = f"{text}\n\n{extras[i]}".strip()
        pages.append((i + 1, text))
    result = _pages_to_chunks(pages)
    if result:
        return result
    if transcribe is not None:
        # Per-page reading found nothing (or the file couldn't be split): read it whole.
        return _scanned_pdf(content, transcribe)
    return _need(result, "PDF")


def _docx(content: bytes, transcribe: Transcriber = None) -> list[dict]:
    from docx import Document
    from docx.oxml.ns import qn
    try:
        doc = Document(io.BytesIO(content))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        # Text boxes and shapes (w:txbxContent) aren't in doc.paragraphs.
        for box in doc.element.body.iter(qn("w:txbxContent")):
            parts.append(" ".join(t.text or "" for t in box.iter(qn("w:t"))))
        seen: set[str] = set()
        for section in doc.sections:
            for part in (section.header, section.footer):
                text = "\n".join(p.text for p in part.paragraphs if p.text.strip())
                if text and text not in seen:
                    seen.add(text)
                    parts.append(text)
        images = [(rel.target_part.blob, rel.target_part.content_type)
                  for rel in doc.part.rels.values() if "image" in rel.reltype and not rel.is_external]
    except Exception as exc:
        log.warning("DOCX parse failed: %s", exc)
        raise HTTPException(status_code=400, detail=(
            "Could not parse Word document: the file may be corrupted. Legacy .doc isn't supported; save as .docx."))
    parts += [f"[Image] {t}" for t in _read_images(images, transcribe)]
    full_text = "\n".join(t for t in parts if t and t.strip())
    return _need(_pages_to_chunks([(1, full_text)]), "Word document")


def _rows_to_chunks(sheet_no: int, sheet_name: str, rows: list[list[str]]) -> list[dict]:
    """Each row keeps its column headers ("Item: Crown | Price: 450") so a single
    chunk still makes sense on its own."""
    rows = [[(c or "").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:] or []
    if not body:                       # a one-row sheet: no header to pair with
        header, body = [], rows
    lines = []
    for r in body:
        if header:
            cells = [f"{h or f'Column {i + 1}'}: {v}" for i, (h, v) in enumerate(zip(header, r)) if v]
        else:
            cells = [v for v in r if v]
        if cells:
            lines.append(" | ".join(cells))
    out = []
    title = f"Sheet: {sheet_name}\n" if sheet_name else ""
    for idx in range(0, len(lines), ROWS_PER_CHUNK):
        out.append({"page": sheet_no, "chunk_index": idx // ROWS_PER_CHUNK,
                    "text": title + "\n".join(lines[idx:idx + ROWS_PER_CHUNK])})
    return out


def _xlsx(content: bytes) -> list[dict]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        log.warning("XLSX parse failed: %s", exc)
        raise HTTPException(status_code=400, detail="Could not open this Excel file. Save it as .xlsx and try again.")
    out: list[dict] = []
    for n, ws in enumerate(wb.worksheets, start=1):
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        out += _rows_to_chunks(n, ws.title, rows)
    return _need(out, "spreadsheet")


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _csv(content: bytes) -> list[dict]:
    text = _decode(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return _need(_rows_to_chunks(1, "", rows), "CSV file")


def _pptx(content: bytes, transcribe: Transcriber = None) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    try:
        deck = Presentation(io.BytesIO(content))
    except Exception as exc:
        log.warning("PPTX parse failed: %s", exc)
        raise HTTPException(status_code=400, detail="Could not open this PowerPoint file. Save it as .pptx and try again.")

    def walk(shapes, parts: list[str], images: list[tuple[bytes, str]]) -> None:
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, parts, images)
                continue
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    images.append((shape.image.blob, shape.image.content_type))
                except Exception:
                    pass

    pages = []
    for n, slide in enumerate(deck.slides, start=1):
        parts: list[str] = []
        images: list[tuple[bytes, str]] = []
        walk(slide.shapes, parts, images)
        parts += [f"[Image] {t}" for t in _read_images(images, transcribe)]
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        pages.append((n, "\n".join(p for p in parts if p.strip())))
    return _need(_pages_to_chunks(pages), "presentation")


class _TextOnly(HTMLParser):
    SKIP = {"script", "style", "noscript", "svg", "nav", "footer", "header", "form"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.title = ""
        self._skip = 0
        self._in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        if tag == "title":
            self._in_title = True
        if tag in {"p", "br", "li", "h1", "h2", "h3", "h4", "tr", "div", "section"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1
        if tag == "title":
            self._in_title = False

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self._skip:
            self.parts.append(data)


def html_to_text(html: str) -> tuple[str, str]:
    p = _TextOnly()
    p.feed(html)
    text = re.sub(r"[ \t]+", " ", "".join(p.parts))
    text = re.sub(r"\n\s*\n+", "\n\n", text).strip()
    return p.title.strip(), text


def _html(content: bytes) -> list[dict]:
    _, text = html_to_text(_decode(content))
    return _need(_pages_to_chunks([(1, text)]), "web page")


def _text(content: bytes) -> list[dict]:
    return _need(_pages_to_chunks([(1, _decode(content))]), "text file")


def extract_chunks(content: bytes, file_name: str, transcribe: Transcriber = None) -> list[dict]:
    ext = _ext(file_name)
    if content[:4] == PDF_MAGIC:
        return _pdf(content, transcribe)
    if content[:4] == ZIP_MAGIC:
        if ext == "docx":
            return _docx(content, transcribe)
        if ext == "xlsx":
            return _xlsx(content)
        if ext == "pptx":
            return _pptx(content, transcribe)
    elif ext == "csv":
        return _csv(content)
    elif ext in TEXT_TYPES:
        return _text(content)
    elif ext in HTML_TYPES:
        return _html(content)
    elif ext in IMAGE_TYPES:
        text = _transcribe(content, IMAGE_TYPES[ext], transcribe, "photo")
        return _need(_pages_to_chunks([(1, text)]), "photo")
    elif ext in AUDIO_TYPES:
        text = _transcribe(content, AUDIO_TYPES[ext], transcribe, "recording")
        return _need(_pages_to_chunks([(1, text)]), "recording")
    raise HTTPException(status_code=400, detail=f"Unsupported file type. Upload {SUPPORTED_LABEL}.")


# ── Web page import (with SSRF protection) ───────────────────────────────────
MAX_PAGE_BYTES = 5 * 1024 * 1024


def _public_host(host: str) -> bool:
    """True only if every address the host resolves to is public internet: blocks
    localhost, private ranges and the cloud metadata server (169.254.169.254)."""
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if (ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved
                or ip.is_multicast or ip.is_unspecified):
            return False
    return True


def fetch_raw(url: str, opener=None, accept: str = "text/html,*/*",
              allow=("html", "text", "xml")) -> tuple[str, str, bytes]:
    """(final_url, content_type, body) for a public address. Follows up to 3
    redirects, re-checking every hop against private and metadata addresses."""
    import urllib.request

    class _NoRedirect(urllib.request.HTTPRedirectHandler):
        def redirect_request(self, *a, **k):
            return None

    opener = opener or urllib.request.build_opener(_NoRedirect)
    for _ in range(4):
        parts = urllib.parse.urlsplit(url)
        if parts.scheme not in ("http", "https") or not parts.hostname:
            raise HTTPException(status_code=400, detail="Enter a full web address starting with https://")
        if not _public_host(parts.hostname):
            raise HTTPException(status_code=400, detail="That address isn't a public web page.")
        req = urllib.request.Request(url, headers={"User-Agent": "RAGaaS-importer/1.0", "Accept": accept})
        try:
            resp = opener.open(req, timeout=15)
        except urllib.error.HTTPError as e:
            if e.code in (301, 302, 303, 307, 308) and e.headers.get("Location"):
                url = urllib.parse.urljoin(url, e.headers["Location"])
                continue
            raise HTTPException(status_code=400, detail=f"The page answered with an error ({e.code}).")
        except Exception as exc:
            log.info("fetch failed %s: %s", url, exc)
            raise HTTPException(status_code=400, detail="Couldn't reach that page.")
        ctype = resp.headers.get("Content-Type", "")
        if not any(a in ctype for a in allow):
            raise HTTPException(status_code=400, detail="That address isn't a web page. Upload files directly instead.")
        body = resp.read(MAX_PAGE_BYTES + 1)
        if len(body) > MAX_PAGE_BYTES:
            raise HTTPException(status_code=413, detail="That page is too large to import.")
        return url, ctype, body
    raise HTTPException(status_code=400, detail="That page redirects too many times.")


def fetch_page(url: str, opener=None) -> tuple[str, str, str]:
    """(final_url, title, text) for a public web page."""
    final, _ctype, body = fetch_raw(url, opener, allow=("html", "text"))
    title, text = html_to_text(_decode(body))
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text on that page.")
    return final, title, text
