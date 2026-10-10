"""Document picker: per-document profiles and chat scoped to picked documents."""
import io
import os

os.environ["RAGAAS_USE_MEMORY_STORE"] = "1"

from docx import Document
from fastapi.testclient import TestClient

from backend.doc_profile import MAX_QUESTIONS, MAX_SUMMARY, build_profile, extractive_profile
from backend.main import app, insights_store, load_index, save_index, usage_store

client = TestClient(app)
TENANT = "tenant-a"
HEADERS = {"Authorization": "Bearer tenant-a-token"}
VIEWER = {"Authorization": "Bearer demo-sales-token"}   # dev viewer on tenant-demo

DOCS = [
    {
        "tenant_id": TENANT, "file_name": "sales_playbook.pdf", "uploaded_at": "2026-10-01T00:00:00Z",
        "chunks": [
            {"page": 1, "chunk_index": 0, "text": "This playbook explains how the sales team prices annual plans. "
                                                  "Discounts above fifteen percent need approval from the sales director."},
            {"page": 4, "chunk_index": 0, "text": "Every discount approval is recorded in the CRM before a quote is sent."},
        ],
    },
    {
        "tenant_id": TENANT, "file_name": "company_handbook.pdf", "uploaded_at": "2026-10-01T00:00:00Z",
        "chunks": [
            {"page": 11, "chunk_index": 0, "text": "The staff wifi password rotates every Friday and is posted at reception. "
                                                   "Discounts on the cafeteria menu are available to all staff."},
        ],
    },
]


def setup_function() -> None:
    usage_store.reset()
    insights_store.reset()
    idx = load_index()
    idx["documents"] = [d for d in idx["documents"] if d["tenant_id"] != TENANT]
    idx["documents"].extend([{**d, "chunks": [dict(c) for c in d["chunks"]]} for d in DOCS])
    save_index(idx)


# ── profiles ──────────────────────────────────────────────────────────────────

def test_extractive_profile_shape() -> None:
    p = extractive_profile("sales_playbook.pdf", DOCS[0]["chunks"])
    assert p["title"] == "Sales playbook"
    assert p["summary"].startswith("This playbook explains")
    assert len(p["summary"]) <= MAX_SUMMARY + 1
    assert 0 < len(p["questions"]) <= MAX_QUESTIONS
    assert all(q.endswith("?") for q in p["questions"])
    assert p["profile_source"] == "extractive"


def test_model_profile_is_sanitised_and_falls_back() -> None:
    class Good:
        def describe(self, title, excerpt):
            return {"title": "  Sales   Playbook ", "summary": "Pricing rules.\n\nFor reps.",
                    "questions": ["Who approves discounts", 42, "x?", "When is a quote logged?"] + ["Q number %d?" % i for i in range(9)]}

    class Broken:
        def describe(self, title, excerpt):
            raise RuntimeError("model down")

    good = build_profile("sales_playbook.pdf", DOCS[0]["chunks"], Good())
    assert good["profile_source"] == "generated"
    assert good["title"] == "Sales Playbook"
    assert good["summary"] == "Pricing rules. For reps."
    assert good["questions"][:2] == ["Who approves discounts?", "When is a quote logged?"]
    assert len(good["questions"]) == MAX_QUESTIONS
    assert good["pages"] == 4

    broken = build_profile("sales_playbook.pdf", DOCS[0]["chunks"], Broken())
    assert broken["profile_source"] == "extractive"
    assert broken["summary"]


def test_listing_backfills_and_persists_profiles() -> None:
    res = client.get("/api/documents", headers=HEADERS)
    assert res.status_code == 200
    by_name = {d["file_name"]: d for d in res.json()}
    playbook = by_name["sales_playbook.pdf"]
    assert playbook["title"] == "Sales playbook"
    assert playbook["summary"] and playbook["questions"]
    assert playbook["pages"] == 4
    stored = next(d for d in load_index()["documents"]
                  if d["tenant_id"] == TENANT and d["file_name"] == "sales_playbook.pdf")
    assert stored["summary"] == playbook["summary"]       # written back, not recomputed each time
    assert len(stored["chunks"]) == 2                      # chunks untouched by the meta update


def test_outdated_extractive_profiles_are_rebuilt_but_model_ones_kept() -> None:
    idx = load_index()
    for d in idx["documents"]:
        if d["tenant_id"] == TENANT and d["file_name"] == "sales_playbook.pdf":
            d.update(summary="old heuristic", questions=["Old?"], profile_source="extractive")   # no version
        if d["tenant_id"] == TENANT and d["file_name"] == "company_handbook.pdf":
            d.update(summary="Written by the model.", questions=["Model question?"], profile_source="generated")
    save_index(idx)
    by_name = {d["file_name"]: d for d in client.get("/api/documents", headers=HEADERS).json()}
    assert by_name["sales_playbook.pdf"]["summary"] != "old heuristic"
    assert by_name["company_handbook.pdf"]["summary"] == "Written by the model."


def test_extractive_profile_skips_header_lines_and_keeps_acronyms() -> None:
    chunks = [{"page": 1, "chunk_index": 0, "text":
               "IT Security Policy Version 3.1 | Effective: March 1, 2025 | Owner: IT Security Team "
               "Password Policy: all employee passwords must meet the following requirements. "
               "This policy explains how staff protect company laptops, accounts and customer data. "
               "Laptop Encryption: every company laptop must use full disk encryption at all times."}]
    p = extractive_profile("it_security_policy.pdf", chunks)
    assert p["title"] == "IT Security Policy"
    assert p["summary"].startswith("This policy explains")
    assert "|" not in p["summary"]
    assert any("password policy" in q for q in p["questions"])
    assert not any("effective" in q.lower() for q in p["questions"])


def test_regenerate_profile_requires_uploader_and_existing_doc() -> None:
    res = client.post("/api/documents/sales_playbook.pdf/profile", headers=HEADERS)
    assert res.status_code == 200
    assert res.json()["file_name"] == "sales_playbook.pdf"
    assert client.post("/api/documents/nope.pdf/profile", headers=HEADERS).status_code == 404
    assert client.post("/api/documents/sales_playbook.pdf/profile", headers=VIEWER).status_code == 403


def test_upload_stores_profile() -> None:
    doc = Document()
    doc.add_paragraph("The travel policy covers flights, hotels and per diem for business trips.")
    doc.add_paragraph("Book flights through the travel desk at least two weeks before departure.")
    buf = io.BytesIO(); doc.save(buf)
    res = client.post("/api/upload", headers=HEADERS,
                      files={"file": ("travel_policy.docx", buf.getvalue(),
                                      "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert res.status_code == 200, res.text
    assert res.json()["title"] == "Travel policy"
    listed = {d["file_name"]: d for d in client.get("/api/documents", headers=HEADERS).json()}
    assert listed["travel_policy.docx"]["summary"].startswith("The travel policy covers")
    client.delete("/api/documents/travel_policy.docx", headers=HEADERS)


# ── scoped chat ───────────────────────────────────────────────────────────────

def test_chat_scoped_to_picked_documents() -> None:
    # "discounts" appears in both documents; unscoped retrieval sees both.
    res = client.post("/api/chat", json={"message": "discounts approval"}, headers=HEADERS)
    assert {c["file_name"] for c in res.json()["citations"]} >= {"sales_playbook.pdf"}

    res = client.post("/api/chat", headers=HEADERS,
                      json={"message": "discounts", "file_names": ["company_handbook.pdf"]})
    body = res.json()
    assert res.status_code == 200
    assert body["citations"] and {c["file_name"] for c in body["citations"]} == {"company_handbook.pdf"}
    assert body["retrieval"]["scoped_to"] == ["company_handbook.pdf"]


def test_chat_scope_drops_unknown_names_and_rejects_empty_scope() -> None:
    res = client.post("/api/chat", headers=HEADERS,
                      json={"message": "wifi password", "file_names": ["company_handbook.pdf", "deleted.pdf"]})
    assert res.status_code == 200
    assert res.json()["retrieval"]["scoped_to"] == ["company_handbook.pdf"]

    used_before = usage_store.get_count(TENANT)
    res = client.post("/api/chat", headers=HEADERS, json={"message": "wifi", "file_names": ["deleted.pdf"]})
    assert res.status_code == 400
    assert usage_store.get_count(TENANT) == used_before   # a bad pick doesn't spend quota


def test_other_tenants_documents_cannot_be_picked() -> None:
    res = client.post("/api/chat", headers={"Authorization": "Bearer tenant-b-token"},
                      json={"message": "wifi", "file_names": ["company_handbook.pdf"]})
    assert res.status_code == 400


def test_scoped_miss_is_not_a_knowledge_gap() -> None:
    q = "what is the parking policy for visitors"
    client.post("/api/chat", headers=HEADERS, json={"message": q, "file_names": ["sales_playbook.pdf"]})
    gaps = client.get("/api/insights", headers=HEADERS).json()["gaps"]
    assert q not in [g["question"] for g in gaps]

    client.post("/api/chat", headers=HEADERS, json={"message": q})
    gaps = client.get("/api/insights", headers=HEADERS).json()["gaps"]
    assert q in [g["question"] for g in gaps]
