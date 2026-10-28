"""Nothing readable is lost to how a file was made: pages with images, math fonts,
garbled or missing text are read by the transcriber page by page; form fields,
comments and images inside Word/PowerPoint files are indexed too."""
import io

import pytest
from fpdf import FPDF
from PIL import Image
from pypdf import PdfReader, PdfWriter
from pypdf.annotations import FreeText
from pypdf.generic import DictionaryObject, NameObject

from backend import ingest

LONG = ("The leave policy gives every employee twenty four days of paid leave each year. "
        "Requests go through the HR portal two weeks ahead. ") * 3


class Reader:
    """Stand-in transcriber that records what it was asked to read."""

    def __init__(self, reply="Read by the transcriber: handwritten note says call Priya.", fail=False):
        self.calls: list[tuple[str, int]] = []
        self.reply, self.fail = reply, fail

    def __call__(self, data: bytes, mime: str) -> str:
        self.calls.append((mime, len(data)))
        if self.fail:
            raise RuntimeError("model unavailable")
        return self.reply


def png(side=200) -> bytes:
    buf = io.BytesIO()
    Image.effect_noise((side, side), 60).convert("RGB").save(buf, "PNG")
    return buf.getvalue()


def pdf(pages: list[dict]) -> bytes:
    doc = FPDF()
    doc.set_font("Helvetica", size=11)
    for p in pages:
        doc.add_page()
        if p.get("text"):
            doc.multi_cell(0, 6, p["text"])
        if p.get("image"):
            doc.image(io.BytesIO(png()), x=20, y=120, w=60)
    return bytes(doc.output())


def texts(chunks):
    return {c["page"]: " ".join(x["text"] for x in chunks if x["page"] == c["page"]) for c in chunks}


# ── PDF pages ─────────────────────────────────────────────────────────────────
def test_plain_text_pdf_is_not_sent_to_the_reader():
    reader = Reader()
    chunks = ingest.extract_chunks(pdf([{"text": LONG}]), "policy.pdf", reader)
    assert reader.calls == [] and "twenty four days" in chunks[0]["text"]


def test_mixed_pdf_reads_only_the_scanned_page():
    reader = Reader(reply="Scanned page: signed by the director on 3 March.")
    chunks = ingest.extract_chunks(pdf([{"text": LONG}, {}]), "mixed.pdf", reader)
    assert len(reader.calls) == 1 and reader.calls[0][0] == "application/pdf"
    by_page = texts(chunks)
    assert "twenty four days" in by_page[1] and "signed by the director" in by_page[2]


def test_page_with_an_image_is_read_and_keeps_its_text():
    reader = Reader()
    chunks = ingest.extract_chunks(pdf([{"text": LONG, "image": True}]), "memo.pdf", reader)
    assert len(reader.calls) == 1
    text = texts(chunks)[1]
    assert "twenty four days" in text and "handwritten note says call Priya" in text   # short read is appended


def test_a_full_read_replaces_a_mangled_text_layer():
    full = LONG + " Constraints: $1 \\le W[i] \\le 10^9$."
    chunks = ingest.extract_chunks(pdf([{"text": LONG, "image": True}]), "task.pdf", Reader(reply=full))
    text = texts(chunks)[1]
    assert "10^9" in text and "[Also on this page]" not in text


def test_math_fonts_flag_a_page():
    page = PdfReader(io.BytesIO(pdf([{"text": LONG}]))).pages[0]
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/BaseFont"): NameObject("/DAAAAA+KaTeX_Math-Italic")})
    page["/Resources"]["/Font"][NameObject("/F9")] = font
    assert "math or symbol fonts" in ingest._pdf_page_reasons(page, LONG)


def test_garbled_text_flags_a_page():
    page = PdfReader(io.BytesIO(pdf([{"text": LONG}]))).pages[0]
    assert "garbled text" in ingest._pdf_page_reasons(page, "(cid:12)(cid:40)(cid:7)(cid:9) " * 20)


def test_one_unreadable_page_does_not_fail_the_upload():
    chunks = ingest.extract_chunks(pdf([{"text": LONG, "image": True}]), "memo.pdf", Reader(fail=True))
    assert "twenty four days" in chunks[0]["text"]


def test_without_a_reader_text_pages_still_index():
    chunks = ingest.extract_chunks(pdf([{"text": LONG, "image": True}]), "memo.pdf", None)
    assert "twenty four days" in chunks[0]["text"]


def test_fully_scanned_pdf_is_read_page_by_page():
    reader = Reader(reply="=== Page 1 ===\nPage content from the scan.")
    chunks = ingest.extract_chunks(pdf([{}, {}, {}]), "scan.pdf", reader)
    assert len(reader.calls) == 3
    assert sorted(texts(chunks)) == [1, 2, 3] and "=== Page" not in chunks[0]["text"]


def test_comments_and_form_values_are_indexed():
    w = PdfWriter(clone_from=io.BytesIO(pdf([{"text": LONG}])))
    w.add_annotation(0, FreeText(text="Reviewer: check the leave cap", rect=(50, 50, 300, 100)))
    buf = io.BytesIO()
    w.write(buf)
    chunks = ingest.extract_chunks(buf.getvalue(), "reviewed.pdf", Reader())
    assert "check the leave cap" in " ".join(c["text"] for c in chunks)


# ── Word and PowerPoint ───────────────────────────────────────────────────────
def test_docx_headers_footers_and_images_are_indexed():
    from docx import Document
    d = Document()
    d.add_paragraph(LONG)
    d.sections[0].header.paragraphs[0].text = "Confidential: HR use only"
    d.sections[0].footer.paragraphs[0].text = "Policy HR-12 revision 4"
    d.add_picture(io.BytesIO(png()))
    buf = io.BytesIO()
    d.save(buf)
    reader = Reader(reply="Org chart: Priya reports to the COO.")
    text = ingest.extract_chunks(buf.getvalue(), "policy.docx", reader)[0]["text"]
    assert "Confidential: HR use only" in text and "Policy HR-12 revision 4" in text
    assert "[Image] Org chart: Priya reports to the COO." in text and reader.calls[0][0] == "image/png"


def test_pptx_pictures_inside_groups_are_read():
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Quarterly targets"
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(png()), Inches(1), Inches(2))
    buf = io.BytesIO()
    deck.save(buf)
    reader = Reader(reply="Whiteboard photo: target 40 new clients.")
    text = ingest.extract_chunks(buf.getvalue(), "deck.pptx", reader)[0]["text"]
    assert "Quarterly targets" in text and "target 40 new clients" in text


def test_tiny_embedded_images_are_skipped():
    reader = Reader()
    assert ingest._read_images([(b"\x89PNG" + b"0" * 100, "image/png"), (png(), "image/gif")], reader) == []
    assert reader.calls == []


@pytest.mark.parametrize("mime", ["image/png", "image/jpeg"])
def test_supported_embedded_images_are_read(mime):
    reader = Reader()
    assert ingest._read_images([(png(), mime)], reader) == [reader.reply]
