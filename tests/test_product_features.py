"""New product surface: file types, web import, feedback, streaming, saved
conversations, gap answers, richer insights and the weekly digest."""
import io
import json
import os

os.environ["RAGAAS_USE_MEMORY_STORE"] = "1"

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

import backend.main as main
from backend import ingest
from backend.conversations import MemoryConversationStore
from backend.digest import MemoryDigestStore
from backend.insights import _aggregate

client = TestClient(main.app)
A = {"Authorization": "Bearer tenant-a-token"}
B = {"Authorization": "Bearer tenant-b-token"}
VIEWER = {"Authorization": "Bearer demo-sales-token"}


@pytest.fixture(autouse=True)
def fresh(monkeypatch):
    main.usage_store.reset()
    main.insights_store.reset()
    monkeypatch.setattr(main, "conversation_store", MemoryConversationStore())
    monkeypatch.setattr(main, "digest_store", MemoryDigestStore())
    idx = main.load_index()
    idx["documents"] = [d for d in idx["documents"] if d["tenant_id"] not in ("tenant-a", "tenant-b")]
    main.save_index(idx)


def upload(name, content, headers=A):
    return client.post("/api/upload", headers=headers, files={"file": (name, content, "application/octet-stream")})


def seed_handbook():
    r = upload("handbook.txt", b"Security awareness training is due within five business days of starting.")
    assert r.status_code == 200, r.text


# ── File types ────────────────────────────────────────────────────────────────
def test_csv_rows_keep_their_headers():
    r = upload("prices.csv", b"Item,Price\nCrown,450\nCleaning,80\n")
    assert r.status_code == 200, r.text
    chunks = main.index_store.doc_chunks("tenant-a", "prices.csv", 10)
    assert "Item: Crown | Price: 450" in chunks[0]["text"]


def test_xlsx_each_sheet_is_a_page():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Rota"
    ws.append(["Day", "Opens"])
    ws.append(["Saturday", "10am"])
    wb.create_sheet("Prices").append(["Item", "Price"])
    wb["Prices"].append(["Crown", 450])
    buf = io.BytesIO()
    wb.save(buf)
    r = upload("ops.xlsx", buf.getvalue())
    assert r.status_code == 200, r.text
    texts = {c["page"]: c["text"] for c in main.index_store.doc_chunks("tenant-a", "ops.xlsx", 10)}
    assert "Day: Saturday | Opens: 10am" in texts[1] and "Sheet: Prices" in texts[2]


def test_pptx_slides_and_notes():
    from pptx import Presentation
    deck = Presentation()
    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = "Fire drill"
    s.placeholders[1].text = "Meet at the car park"
    s.notes_slide.notes_text_frame.text = "Warden counts heads"
    buf = io.BytesIO()
    deck.save(buf)
    r = upload("training.pptx", buf.getvalue())
    assert r.status_code == 200, r.text
    text = main.index_store.doc_chunks("tenant-a", "training.pptx", 5)[0]["text"]
    assert "car park" in text and "Warden" in text


def test_markdown_html_text():
    assert upload("faq.md", "# Hours\nWe open at nine.".encode()).status_code == 200
    html = b"<html><title>Returns</title><script>x()</script><p>Returns within 30 days.</p></html>"
    assert upload("returns.html", html).status_code == 200
    text = main.index_store.doc_chunks("tenant-a", "returns.html", 5)[0]["text"]
    assert "30 days" in text and "x()" not in text


def test_photo_and_audio_are_transcribed():
    assert upload("notice.jpg", b"\xff\xd8\xff fake jpeg").status_code == 200
    assert upload("memo.m4a", b"fake audio").status_code == 200
    assert "Dev transcript of a recording" in main.index_store.doc_chunks("tenant-a", "memo.m4a", 1)[0]["text"]


def test_scanned_pdf_falls_back_to_reading_it():
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()                              # blank page: no selectable text
    r = upload("scan.pdf", bytes(pdf.output()))
    assert r.status_code == 200, r.text
    assert "Dev transcript of a scan" in main.index_store.doc_chunks("tenant-a", "scan.pdf", 1)[0]["text"]


def test_unsupported_type_lists_what_works():
    r = upload("tool.exe", b"MZ\x90\x00")
    assert r.status_code == 400 and "Excel" in r.json()["detail"]


def test_scanned_page_markers_split_pages():
    chunks = ingest._scanned_pdf(b"%PDF", lambda d, m: "=== Page 1 ===\nfirst\n=== Page 2 ===\nsecond")
    assert [(c["page"], c["text"]) for c in chunks] == [(1, "first"), (2, "second")]


# ── Web import ────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("url", ["http://127.0.0.1/admin", "http://169.254.169.254/computeMetadata/v1/",
                                 "http://localhost:8000/", "ftp://example.com/x", "file:///etc/passwd"])
def test_import_refuses_private_and_non_web_addresses(url):
    r = client.post("/api/import-url", headers=A, json={"url": url})
    assert r.status_code == 400


def test_import_url_indexes_the_page(monkeypatch):
    monkeypatch.setattr(main, "fetch_page", lambda url: (url, "Refund policy", "Refunds are issued within 14 days."))
    r = client.post("/api/import-url", headers=A, json={"url": "https://shop.example.com/refunds"})
    assert r.status_code == 200, r.text
    assert r.json()["file_name"].endswith(".html")
    assert client.post("/api/import-url", headers=VIEWER, json={"url": "https://x.example.com"}).status_code == 403


# ── Feedback, streaming, conversations ───────────────────────────────────────
def test_feedback_on_an_answer():
    seed_handbook()
    body = client.post("/api/chat", headers=A, json={"message": "when is security training due"}).json()
    assert body["answer_id"] and body["conversation_id"]
    assert client.post("/api/feedback", headers=A, json={"answer_id": body["answer_id"], "helpful": False}).status_code == 200
    assert client.post("/api/feedback", headers=A, json={"answer_id": "nope", "helpful": True}).status_code == 404
    fb = client.get("/api/insights", headers=A).json()["feedback"]
    assert fb["unhelpful"] == 1 and fb["recent_unhelpful"][0]["question"] == "when is security training due"


def test_stream_sends_sources_then_text_then_done():
    seed_handbook()
    r = client.post("/api/chat/stream", headers=A, json={"message": "when is security training due"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/x-ndjson")
    events = [json.loads(l) for l in r.text.splitlines() if l.strip()]
    assert events[0]["type"] == "meta" and events[0]["citations"]
    deltas = "".join(e["text"] for e in events if e["type"] == "delta")
    done = events[-1]
    assert done["type"] == "done" and done["answer"] == deltas.strip()
    conv = client.get(f"/api/conversations/{done['conversation_id']}", headers=A).json()
    assert [m["role"] for m in conv["messages"]] == ["user", "assistant"]
    assert conv["messages"][1]["answer_id"] == done["answer_id"]


def test_stream_respects_the_daily_limit(monkeypatch):
    def full(_t):
        raise HTTPException(status_code=429, detail="Tenant query quota exceeded")
    monkeypatch.setattr(main, "enforce_quota", full)
    assert client.post("/api/chat/stream", headers=A, json={"message": "hi there"}).status_code == 429


def test_conversations_append_list_and_stay_private():
    seed_handbook()
    first = client.post("/api/chat", headers=A, json={"message": "when is security training due"}).json()
    cid = first["conversation_id"]
    client.post("/api/chat", headers=A, json={"message": "and who runs it?", "conversation_id": cid})
    listed = client.get("/api/conversations", headers=A).json()
    assert listed[0]["id"] == cid and listed[0]["messages"] == 4
    assert client.get(f"/api/conversations/{cid}", headers=B).status_code == 404
    assert client.delete(f"/api/conversations/{cid}", headers=A).status_code == 200
    assert client.get("/api/conversations", headers=A).json() == []


# ── Gap answers ───────────────────────────────────────────────────────────────
def test_written_gap_answer_becomes_searchable_and_replaces_itself():
    payload = {"question": "Can patients pay in instalments?", "answer": "Yes, up to three monthly payments over $600."}
    assert client.post("/api/gaps/answer", headers=VIEWER, json=payload).status_code == 403
    assert client.post("/api/gaps/answer", headers=A, json=payload).json()["entries"] == 1
    payload["answer"] = "Yes, up to four monthly payments over $600."
    assert client.post("/api/gaps/answer", headers=A, json=payload).json()["entries"] == 1
    r = client.post("/api/chat", headers=A, json={"message": "can patients pay in monthly instalments"}).json()
    assert r["citations"][0]["file_name"] == "team-answers.md"
    assert "four monthly payments" in r["citations"][0]["excerpt"]


# ── Insights ──────────────────────────────────────────────────────────────────
def test_similar_questions_are_grouped():
    ev = [{"question": q, "score": 0.1, "ts": "2026-10-19T10:00:00+00:00"} for q in
          ["Can I pay in instalments?", "can i pay in instalments", "Can patients pay in instalments?",
           "What is the wifi password?"]]
    gaps = _aggregate(ev)["gaps"]
    assert gaps[0]["count"] == 3 and gaps[0]["variants"]
    assert len(gaps) == 2


def test_insights_has_trend_channels_and_unused_documents():
    seed_handbook()
    upload("menu.txt", b"Lunch menu changes on Mondays.")
    client.post("/api/chat", headers=A, json={"message": "when is security training due"})
    data = client.get("/api/insights", headers=A).json()
    assert len(data["trend"]) == 8 and data["trend"][-1]["questions"] == 1
    assert data["channels"] == [{"source": "chat", "count": 1}]
    usage = {d["file_name"]: d["citations"] for d in data["documents"]}
    assert usage["handbook.txt"] == 1 and usage["menu.txt"] == 0
    assert "faqs" in data


# ── Digest ────────────────────────────────────────────────────────────────────
def test_digest_preview_subscribe_and_cron(monkeypatch):
    seed_handbook()
    client.post("/api/chat", headers=A, json={"message": "what is the refund window for crowns"})
    prev = client.get("/api/insights/digest", headers=A).json()
    assert "question" in prev["text"] and "this week" in prev["subject"]
    assert client.post("/api/insights/digest/subscribe", headers=A, json={"on": True}).json()["subscribed"]

    sent = []
    monkeypatch.setattr(main, "send_email", lambda to, subj, body: sent.append(to) or True)
    assert client.post("/api/cron/digest").status_code == 403
    monkeypatch.setenv("CRON_TOKEN", "t0k")
    assert client.post("/api/cron/digest", headers={"X-Cron-Token": "wrong"}).status_code == 403
    assert client.post("/api/cron/digest", headers={"X-Cron-Token": "t0k"}).json()["sent"] == 1
    assert sent == ["tenant-a@ragaas.local"]


def test_unexpected_error_is_json_with_cors(monkeypatch):
    """A crash must reach the browser as a readable 500, not a CORS "network error"."""
    def boom(_tenant):
        raise RuntimeError("BulkWriter is closed")
    monkeypatch.setattr(main.index_store, "list_docs", boom)
    r = client.get("/api/documents", headers={**A, "Origin": "http://localhost:5173"})
    assert r.status_code == 500
    assert r.headers.get("access-control-allow-origin") == "http://localhost:5173"
    assert "try again" in r.json()["detail"]